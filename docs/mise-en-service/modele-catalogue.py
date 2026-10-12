# Le modèle de catalogue de Maymar (docs/MISE-EN-SERVICE.md) : la feuille « Catalogue »
# est celle que l'import de SkanEcom lit (la première, en-têtes en ligne 1) ; « Exemple »
# et « Mode d'emploi » ne sont pas lus. python3 modele-catalogue.py modele-catalogue-maymar.xlsx
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1]
LIGNES = 600  # lignes préparées (formats, listes) sous les en-têtes

POLICE = "Arial"
ENCRE = "1F2328"
TETE_FOND = "1F2328"
OBLIGATOIRE_FOND = "F6E7C1"
SAISIE_FOND = "FFFDF5"
TRAIT = Side(style="thin", color="D0D4DA")

# (en-tête exact lu par l'import, largeur, obligatoire, format, bulle d'aide)
COLONNES = [
    ("Produit", 34, True, "@", "OBLIGATOIRE. Le nom du produit, tel que le client le lira.\nRépétez-le sur chaque ligne de ses variantes (une ligne par couleur ou par taille)."),
    ("Référence", 18, True, "@", "OBLIGATOIRE. Un code unique par ligne (par variante), ex. VAL55-NOIR.\nJamais deux fois le même code dans le fichier."),
    ("Prix (TND)", 13, True, "#,##0.000", "OBLIGATOIRE. Le prix de vente en dinars, TTC.\nEx. 189 ou 189,500."),
    ("Prix barré (TND)", 16, False, "#,##0.000", "Facultatif. L'ancien prix, affiché barré à côté du prix.\nLaissez vide s'il n'y a pas de remise."),
    ("Stock", 9, False, "0", "Facultatif. La quantité disponible de cette variante (un nombre entier).\nVide = on ne change pas le stock."),
    ("Rayon", 26, False, "@", "Le rayon du produit sur la vitrine. Déjà créés : Valises cabine, Valises soute,\nSacs de voyage, Sacs à dos, Accessoires de voyage. Un nouveau nom crée un rayon ;\nun sous-rayon s'écrit avec « > » (Sacs de voyage > Sacs week-end)."),
    ("Couleur", 14, False, "@", "Facultatif. La couleur de cette variante (Noir, Bleu nuit…).\nLaissez vide si le produit n'a qu'une couleur."),
    ("Taille", 18, False, "@", "Facultatif. La taille de cette variante, pour un produit vendu en tailles (S, M, L).\nUne valise de 55 cm et une de 70 cm sont deux produits, chacun dans son rayon."),
    ("Marque", 14, False, "@", "Facultatif. La marque du produit, si elle compte pour le client."),
    ("Description", 44, False, "@", "Facultatif. Quelques phrases sur le produit, sur sa première ligne seulement.\nLe backoffice peut aussi la rédiger plus tard."),
    ("Poids (kg)", 11, False, "0.0", "Facultatif mais utile : le poids du colis en kilos (ex. 2,6).\nIl sert au supplément de livraison selon le poids."),
    ("Volume (L)", 11, False, "0", "Facultatif. La contenance en litres (ex. 40), sur la première ligne du produit.\nAffichée sur la fiche et filtrable sur la vitrine."),
    ("Dimensions", 20, False, "@", "Facultatif. Hauteur × largeur × profondeur, ex. 55 × 40 × 20 cm.\nAffichées sur la fiche."),
    ("Roues", 8, False, "0", "Facultatif. Le nombre de roues (2 ou 4).\nFiltrable sur la vitrine."),
    ("Publié", 9, False, "@", "« oui » : en vente sur la vitrine dès l'import.\n« non » : importé mais caché, à relire d'abord.\nVide = non."),
]

EXEMPLE = [
    ["Valise rigide 4 roues 55 cm", "VAL55-NOIR", 189, None, 12, "Valises cabine", "Noir", "", "", "Coque ABS, quatre roues pivotantes, serrure TSA. Aux dimensions cabine des principales compagnies.", 2.6, 40, "55 × 40 × 20 cm", 4, "oui"],
    ["Valise rigide 4 roues 55 cm", "VAL55-BLEU", 189, None, 4, "Valises cabine", "Bleu nuit", "", "", "", 2.6, None, "", None, "oui"],
    ["Valise rigide 4 roues 70 cm", "VAL70-NOIR", 259, 289, 6, "Valises soute", "Noir", "", "", "La même coque en grande taille, pour deux semaines de voyage.", 3.8, 90, "70 × 46 × 28 cm", 4, "oui"],
    ["Cadenas à code TSA", "CAD-TSA", 25, None, 40, "Accessoires de voyage", "", "", "", "Trois chiffres, accepté par les contrôles des aéroports.", 0.1, None, "", None, "oui"],
]



def police(gras=False, couleur=ENCRE, taille=10, italique=False):
    return Font(name=POLICE, bold=gras, color=couleur, size=taille, italic=italique)


def tete(feuille):
    for i, (nom, largeur, obligatoire, _fmt, aide) in enumerate(COLONNES, start=1):
        c = feuille.cell(row=1, column=i, value=nom)
        c.font = police(gras=True, couleur="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TETE_FOND)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=TRAIT)
        lettre = c.column_letter
        feuille.column_dimensions[lettre].width = largeur
        bulle = Comment(aide, "Maymar")
        bulle.width, bulle.height = 320, 120
        c.comment = bulle
    feuille.row_dimensions[1].height = 30
    feuille.freeze_panes = "B2"


wb = Workbook()

# 1. La feuille lue par l'import --------------------------------------------
cat = wb.active
cat.title = "Catalogue"
tete(cat)
for i, (_nom, _l, obligatoire, fmt, _aide) in enumerate(COLONNES, start=1):
    for r in range(2, LIGNES + 2):
        c = cat.cell(row=r, column=i)
        c.number_format = fmt
        c.font = police()
        c.fill = PatternFill("solid", fgColor=OBLIGATOIRE_FOND if obligatoire else SAISIE_FOND)
        c.border = Border(bottom=TRAIT, right=TRAIT)
        if _nom == "Description":
            c.alignment = Alignment(wrap_text=True, vertical="top")
fin = LIGNES + 1
prix = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                      errorTitle="Prix", error="Un prix en dinars, positif (ex. 189 ou 189,500).")
prix.add(f"C2:D{fin}")
stock = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                       errorTitle="Stock", error="Un nombre entier, 0 ou plus.")
stock.add(f"E2:E{fin}")
poids = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                       errorTitle="Poids", error="Le poids en kilos, entre 0 et 100 (ex. 2,6).")
poids.add(f"K2:K{fin}")
volume = DataValidation(type="decimal", operator="between", formula1="0", formula2="500", allow_blank=True,
                        errorTitle="Volume", error="La contenance en litres (ex. 40).")
volume.add(f"L2:L{fin}")
roues = DataValidation(type="whole", operator="between", formula1="0", formula2="8", allow_blank=True,
                       errorTitle="Roues", error="Le nombre de roues (ex. 4).")
roues.add(f"N2:N{fin}")
publie = DataValidation(type="list", formula1='"oui,non"', allow_blank=True,
                        errorTitle="Publié", error="« oui » ou « non ».")
publie.add(f"O2:O{fin}")
for v in (prix, stock, poids, volume, roues, publie):
    cat.add_data_validation(v)

# 2. L'exemple (non lu) -------------------------------------------------------
ex = wb.create_sheet("Exemple")
tete(ex)
ex.freeze_panes = None
for r, ligne in enumerate(EXEMPLE, start=2):
    for i, v in enumerate(ligne, start=1):
        c = ex.cell(row=r, column=i, value=v if v != "" else None)
        c.number_format = COLONNES[i - 1][3]
        c.font = police()
        c.border = Border(bottom=TRAIT, right=TRAIT)
        if COLONNES[i - 1][0] == "Description":
            c.alignment = Alignment(wrap_text=True, vertical="top")
n = len(EXEMPLE) + 3
ex.cell(row=n, column=1, value="Exemple seulement : cette feuille n'est pas importée. Remplissez la feuille « Catalogue ».").font = police(gras=True, couleur="9A3412")
ex.cell(row=n + 1, column=1, value="Les deux premières lignes sont UN produit (même nom) en deux couleurs : chaque couleur a sa ligne, sa référence, son prix et son stock. La valise de 70 cm est un autre produit, dans un autre rayon.").font = police(italique=True)
ex.cell(row=n + 2, column=1, value="La description, le rayon et les caractéristiques s'écrivent une fois, sur la première ligne du produit ; les lignes suivantes peuvent les laisser vides.").font = police(italique=True)

# 3. Le mode d'emploi -------------------------------------------------------
md = wb.create_sheet("Mode d'emploi")
md.column_dimensions["A"].width = 4
md.column_dimensions["B"].width = 110
lignes_md = [
    ("titre", "Le catalogue de Maymar — mode d'emploi"),
    ("", ""),
    ("gras", "1. Une ligne par variante"),
    ("", "Une valise en noir et en bleu, c'est deux lignes : le même nom de produit, deux références."),
    ("", "Chaque taille de valise est un produit à part (« … 55 cm », « … 70 cm »), rangé dans son rayon : elle a son volume et ses dimensions."),
    ("", "Un produit sans couleur (un cadenas) tient sur une seule ligne : laissez Couleur et Taille vides."),
    ("", ""),
    ("gras", "2. Les colonnes jaunes sont obligatoires"),
    ("", "Produit, Référence, Prix (TND). Tout le reste peut rester vide et se compléter plus tard dans le backoffice."),
    ("", "Passez la souris sur un en-tête (ou touchez-le) pour lire ce qu'il attend."),
    ("", ""),
    ("gras", "3. Les rayons"),
    ("", "Cinq rayons existent déjà : Valises cabine, Valises soute, Sacs de voyage, Sacs à dos, Accessoires de voyage. Reprenez ces noms exactement."),
    ("", "Un autre nom crée un nouveau rayon ; un sous-rayon s'écrit avec « > » : « Sacs de voyage > Sacs week-end ». Trois niveaux au plus."),
    ("", ""),
    ("gras", "4. Les prix et le poids"),
    ("", "En dinars : 189 ou 189,500. Le prix barré est l'ancien prix, seulement s'il y a une remise."),
    ("", "Le poids du colis en kilos (2,6) : il sert au supplément de livraison selon le poids."),
    ("", ""),
    ("gras", "4 bis. Les caractéristiques"),
    ("", "Volume (litres), Dimensions et Roues s'affichent sur la fiche du produit ; le volume et les roues servent de filtres sur la vitrine. Une valeur par produit, sur sa première ligne."),
    ("", ""),
    ("gras", "5. En vente tout de suite, ou à relire d'abord"),
    ("", "Publié = « oui » : le produit est en vente dès l'import. « non » (ou vide) : il est importé caché, à relire dans le backoffice."),
    ("", ""),
    ("gras", "6. Les photos"),
    ("", "Pas de photo dans ce fichier : elles s'envoient à part, dans un dossier (voir le protocole photo)."),
    ("", "Nommez chaque photo d'après la référence de sa ligne : VAL55-NOIR.jpg, puis VAL55-NOIR-2.jpg, VAL55-NOIR-3.jpg… Elles se rangent alors toutes seules sur leur produit et leur couleur."),
    ("", ""),
    ("gras", "7. Ce qui se passe ensuite"),
    ("", "Envoyez le fichier à Skander. L'import montre d'abord un aperçu ligne par ligne : rien n'est écrit tant qu'il y a une erreur, et chaque erreur dit sa ligne."),
    ("", "Le fichier peut être renvoyé plus tard avec seulement les prix et le stock à jour : une case vide ne remplace rien."),
    ("", ""),
    ("note", "Ne changez pas les noms des colonnes de la feuille « Catalogue » et ne mettez rien au-dessus de la ligne des en-têtes : l'import les lit tels quels."),
]
for r, (genre, texte) in enumerate(lignes_md, start=1):
    c = md.cell(row=r, column=2, value=texte or None)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if genre == "titre":
        c.font = police(gras=True, taille=14)
        md.row_dimensions[r].height = 24
    elif genre == "gras":
        c.font = police(gras=True, taille=11)
    elif genre == "note":
        c.font = police(gras=True, couleur="9A3412")
    else:
        c.font = police()
md.sheet_view.showGridLines = False

wb.save(SORTIE)
print("écrit", SORTIE)
